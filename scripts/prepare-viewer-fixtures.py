#!/usr/bin/env python3
"""Distinct real multi-page inputs; test assets never ship in the application."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, re
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from openpyxl import Workbook

root=Path(__file__).resolve().parent.parent
out=root/'artifacts/open-fixtures/android-assets';out.mkdir(parents=True,exist_ok=True)
source=root/'core/build/visual-fixtures/fixture.hwpx'
with ZipFile(source) as z:
    xml=z.read('Contents/section0.xml').decode()
    # Preserve namespace declarations, table and image binary entries byte-for-byte.
    matches=list(re.finditer(r'<hp:p\b[^>]*>.*?</hp:p>',xml,re.S))
    first=next(m for m in matches if 'PAGE_EDIT_ORIGINAL' in m.group(0) and '<hp:tbl' not in m.group(0))
    template=first.group(0)
    paragraphs=[]
    for i in range(20):
        p=template.replace('PAGE_EDIT_ORIGINAL',f'VIEWER_HWPX_PAGE_{i+1:02d}')
        p=re.sub(r'\bpageBreak="[^"]*"',f'pageBreak="{int(i>0)}"',p,count=1)
        assert f'pageBreak="{int(i>0)}"' in p
        # Let the actual layout engine recompute line positions instead of copying
        # stale per-page layout caches from the seed.
        p=re.sub(r'<hp:linesegarray\b.*?</hp:linesegarray>','',p,flags=re.S|re.I)
        paragraphs.append(p)
    generated=xml[:first.start()]+''.join(paragraphs)+xml[first.end():]
    with ZipFile(out/'viewer-20.hwpx','w') as target:
        for info in z.infolist():
            target.writestr(info,generated.encode() if info.filename=='Contents/section0.xml' else z.read(info.filename))

doc=Document()
for i in range(12):
    if i:doc.add_page_break()
    doc.add_heading(f'VIEWER_DOCX_PAGE_{i+1:02d}',1)
    doc.add_paragraph('다중 페이지 읽기 및 마지막 페이지 보존 검증')
    table=doc.add_table(rows=2, cols=2);table.cell(0,0).text=f'TABLE_{i+1:02d}'
doc.save(out/'viewer-12.docx')

slides=Presentation();slides.slide_width=Inches(10);slides.slide_height=Inches(5.625)
for i in range(30):
    s=slides.slides.add_slide(slides.slide_layouts[6])
    s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text=f'VIEWER_SLIDE_{i+1:02d}'
slides.save(out/'viewer-30.pptx')

book=Workbook()
for i in range(3):
    sheet=book.active if i==0 else book.create_sheet()
    sheet.title=f'Viewer {i+1}'
    for row in range(1,5001):sheet.cell(row,1,row);sheet.cell(row,100,f'CELL_CV{row}_S{i+1}')
    sheet['B1']='=SUM(A1:A5000)'
book.save(out/'viewer-large.xlsx')
(out/'viewer-long.md').write_text('\n\n'.join(f'# 제목 {i}\n\nMD_SECTION_{i:03d}\n\n|A|B|\n|-|-|\n|한글|{i}|' for i in range(100)),encoding='utf-8')
(out/'viewer-long.txt').write_text('\n'.join(f'TXT_LINE_{i:05d} 한글 전체 문서' for i in range(20000)),encoding='utf-8')
manifest={p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in out.glob('viewer-*')}
(out.parent/'viewer-fixtures.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('VIEWER_FIXTURES',json.dumps(manifest,ensure_ascii=False))
