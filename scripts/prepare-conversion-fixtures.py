#!/usr/bin/env python3
"""Synthetic layout fixtures; no user documents or signatures."""
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DI
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Border,Side
root=Path('app/src/androidTest/assets/conversion');root.mkdir(parents=True,exist_ok=True)
image=root/'stamp.png'
im=Image.new('RGB',(160,90),'white');ImageDraw.Draw(im).rectangle((8,8,150,80),fill='blue');im.save(image)
ppt=Presentation();ppt.slide_width=Inches(12);ppt.slide_height=Inches(6.75)
for i in range(10):
 s=ppt.slides.add_slide(ppt.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(235,240,250)
 box=s.shapes.add_textbox(Inches(.5),Inches(.3),Inches(8),Inches(1));box.text_frame.text=f'TEST_SLIDE_{i+1:02d}'
 box.text_frame.paragraphs[0].font.size=Pt(30)
 s.shapes.add_picture(str(image),Inches(.8),Inches(2),width=Inches(2))
 shape=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(8),Inches(2),Inches(2),Inches(2));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(210,40,60)
 table=s.shapes.add_table(2,2,Inches(3.5),Inches(2),Inches(3),Inches(2)).table
 table.cell(0,0).text='TABLE';table.cell(1,1).text=str(i+1)
ppt.save(root/'layout.pptx')
doc=Document();doc.sections[0].header.paragraphs[0].text='HEADER_TEST';doc.sections[0].footer.paragraphs[0].text='FOOTER_TEST'
doc.add_heading('DOCX_PAGE_ONE',0);doc.add_paragraph('Korean 한국어 conversion test');doc.add_table(rows=2,cols=2).cell(0,0).text='TABLE_TEST';doc.add_picture(str(image),width=DI(1));doc.add_page_break();doc.add_paragraph('DOCX_PAGE_TWO');doc.save(root/'layout.docx')
book=Workbook();a=book.active;a.title='First';a['A1']='XLSX_SHEET_ONE';a['A2']=20;a['B2']=22;a['C2']='=A2+B2';a.merge_cells('A3:C3');a['A3']='MERGED_TEST'
a['A1'].fill=PatternFill('solid',fgColor='FFFF00');a['C2'].number_format='0.00';a.sheet_properties.pageSetUpPr.fitToPage=True;a.page_setup.orientation='landscape';a.page_setup.fitToWidth=1;a.page_setup.fitToHeight=1;a.print_area='A1:D10'
b=book.create_sheet('Second');b['A1']='XLSX_SHEET_TWO';b['A2']=123;b.page_setup.fitToWidth=1;b.print_area='A1:D10'
book.save(root/'layout.xlsx')
print('CONVERSION_FIXTURES',[(p.name,p.stat().st_size) for p in root.iterdir()])
